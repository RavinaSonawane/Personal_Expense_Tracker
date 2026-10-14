# =========================================================
#  Tests for Personal Expense & Budget Tracker
# =========================================================

import os
import openpyxl
import pytest
from expense_tracker import (
    init_db,
    add_expense,
    get_all_expenses,
    get_total_expenses,
    get_category_summary,
    delete_expense,
    export_to_excel,
    import_from_excel,
)


@pytest.fixture
def temp_db(tmp_path):
    """Fixture providing a fresh isolated database file for each test."""
    db_file = str(tmp_path / "test_expenses.db")
    init_db(db_file)
    return db_file


def test_init_db(temp_db):
    """Test initializing database and creating table."""
    expenses = get_all_expenses(temp_db)
    assert expenses == []
    assert get_total_expenses(temp_db) == 0.0


def test_add_expense_valid(temp_db):
    """Test adding a valid expense."""
    exp_id = add_expense("Lunch", 15.50, "Food", "2026-10-14", temp_db)
    assert exp_id is not None
    assert exp_id > 0

    all_exp = get_all_expenses(temp_db)
    assert len(all_exp) == 1
    assert all_exp[0]["title"] == "Lunch"
    assert all_exp[0]["amount"] == 15.50
    assert all_exp[0]["category"] == "Food"
    assert all_exp[0]["date"] == "2026-10-14"


def test_add_expense_invalid_amount(temp_db):
    """Test adding expense with negative or zero amount raises ValueError."""
    with pytest.raises(ValueError, match="Expense amount must be greater than zero."):
        add_expense("Free Item", 0, "Other", db_path=temp_db)

    with pytest.raises(ValueError, match="Expense amount must be greater than zero."):
        add_expense("Refund", -10, "Other", db_path=temp_db)


def test_add_expense_invalid_title(temp_db):
    """Test adding expense with empty title raises ValueError."""
    with pytest.raises(ValueError, match="Expense title cannot be empty."):
        add_expense("", 25.0, "Food", db_path=temp_db)

    with pytest.raises(ValueError, match="Expense title cannot be empty."):
        add_expense("   ", 25.0, "Food", db_path=temp_db)


def test_add_expense_invalid_category(temp_db):
    """Test adding expense with empty category raises ValueError."""
    with pytest.raises(ValueError, match="Category cannot be empty."):
        add_expense("Snack", 5.0, "", db_path=temp_db)


def test_get_total_expenses(temp_db):
    """Test calculation of total expenses."""
    add_expense("Coffee", 4.50, "Food", db_path=temp_db)
    add_expense("Bus", 2.50, "Transport", db_path=temp_db)
    add_expense("Book", 13.00, "Education", db_path=temp_db)

    total = get_total_expenses(temp_db)
    assert total == 20.00


def test_get_category_summary(temp_db):
    """Test summary grouped by category."""
    add_expense("Breakfast", 10.00, "Food", db_path=temp_db)
    add_expense("Dinner", 25.00, "Food", db_path=temp_db)
    add_expense("Train", 15.00, "Transport", db_path=temp_db)

    summary = get_category_summary(temp_db)
    assert summary["Food"] == 35.00
    assert summary["Transport"] == 15.00


def test_delete_expense(temp_db):
    """Test deleting an expense by ID."""
    exp_id = add_expense("Unwanted item", 50.0, "Misc", db_path=temp_db)
    assert get_total_expenses(temp_db) == 50.0

    # Delete existing
    deleted = delete_expense(exp_id, temp_db)
    assert deleted is True
    assert get_total_expenses(temp_db) == 0.0

    # Delete non-existent ID
    deleted_again = delete_expense(9999, temp_db)
    assert deleted_again is False


def test_export_to_excel(temp_db, tmp_path):
    """Test exporting SQLite expenses to formatted Excel document with both tabs."""
    add_expense("Groceries", 65.50, "Food", "2026-10-10", temp_db)
    add_expense("Internet", 50.00, "Bills", "2026-10-12", temp_db)

    excel_file = str(tmp_path / "test_export.xlsx")
    saved_path = export_to_excel(temp_db, excel_file)
    assert os.path.exists(saved_path)

    wb = openpyxl.load_workbook(saved_path, data_only=False)
    assert "Daily Expenses" in wb.sheetnames
    assert "Summary & Analytics" in wb.sheetnames

    # Check Daily Expenses tab headers and rows
    ws_daily = wb["Daily Expenses"]
    assert ws_daily.cell(row=1, column=1).value == "ID"
    assert ws_daily.cell(row=1, column=5).value == "Amount ($)"
    assert ws_daily.cell(row=1, column=6).value == "Week Num"

    # Row 2 (first expense)
    assert ws_daily.cell(row=2, column=3).value == "Food"
    assert ws_daily.cell(row=2, column=4).value == "Groceries"
    assert ws_daily.cell(row=2, column=5).value == 65.50
    # Check Week, Month, Year formulas
    assert "ISOWEEKNUM" in str(ws_daily.cell(row=2, column=6).value)
    assert "yyyy-mm" in str(ws_daily.cell(row=2, column=7).value)
    assert "YEAR" in str(ws_daily.cell(row=2, column=8).value)

    # Check Summary & Analytics tab
    ws_sum = wb["Summary & Analytics"]
    assert "SUMMARY & ANALYTICS" in str(ws_sum["A1"].value)
    assert "=SUM('Daily Expenses'!E2:E1000)" == ws_sum["B5"].value
    # Monthly formula check
    assert "=SUMIF('Daily Expenses'!$G$2:$G$1000, A10, 'Daily Expenses'!$E$2:$E$1000)" == ws_sum["C10"].value
    # Weekly formula check
    assert "=SUMIF('Daily Expenses'!$F$2:$F$1000, E10, 'Daily Expenses'!$E$2:$E$1000)" == ws_sum["G10"].value


def test_import_from_excel(temp_db, tmp_path):
    """Test importing new expenses entered in the Excel document back to SQLite."""
    # Export empty or base db
    add_expense("Coffee", 4.50, "Food", "2026-10-10", temp_db)
    excel_file = str(tmp_path / "test_sync.xlsx")
    export_to_excel(temp_db, excel_file)

    # User opens Excel and types a new expense into row 3
    wb = openpyxl.load_workbook(excel_file)
    ws = wb["Daily Expenses"]
    ws.cell(row=3, column=2, value="2026-10-15")
    ws.cell(row=3, column=3, value="Shopping")
    ws.cell(row=3, column=4, value="Running Shoes")
    ws.cell(row=3, column=5, value=89.99)
    wb.save(excel_file)

    # Run import
    imported = import_from_excel(excel_file, temp_db)
    assert imported == 1

    # Verify the new record is saved in SQLite
    all_expenses = get_all_expenses(temp_db)
    assert len(all_expenses) == 2
    titles = [e["title"] for e in all_expenses]
    assert "Running Shoes" in titles
    assert get_total_expenses(temp_db) == round(4.50 + 89.99, 2)

    # Re-importing should not duplicate
    reimported = import_from_excel(excel_file, temp_db)
    assert reimported == 0
    assert len(get_all_expenses(temp_db)) == 2
