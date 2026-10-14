# =========================================================
#  Project: Personal Expense & Budget Tracker (CLI + SQLite + Excel)
#  Author: Ravina Sonawane
#  Description: Manage expenses, track budget, analyze spending,
#               and export/import Excel documents with live weekly,
#               monthly, and yearly calculations.
# =========================================================

import os
import sys
import sqlite3
from datetime import date as dt_date, datetime

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    openpyxl = None


DEFAULT_DB = "expenses.db"
DEFAULT_EXCEL = "expense_tracker.xlsx"

STANDARD_CATEGORIES = [
    "Food",
    "Transport",
    "Bills",
    "Shopping",
    "Entertainment",
    "Health",
    "Education",
    "Other",
]


def get_connection(db_path=DEFAULT_DB):
    """Return a sqlite3 database connection with Row factory."""
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    return conn


def init_db(db_path=DEFAULT_DB):
    """Create the expenses table if it does not already exist."""
    with get_connection(db_path) as conn:
        cursor = conn.cursor()
        cursor.execute(
            """
            CREATE TABLE IF NOT EXISTS expenses (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                title TEXT NOT NULL,
                amount REAL NOT NULL,
                category TEXT NOT NULL,
                date TEXT NOT NULL
            )
            """
        )
        conn.commit()


def add_expense(title, amount, category, expense_date=None, db_path=DEFAULT_DB):
    """
    Add a new expense.
    Validates that title is not empty, amount > 0, and category is provided.
    """
    if not title or not title.strip():
        raise ValueError("Expense title cannot be empty.")

    try:
        amount = float(amount)
    except (TypeError, ValueError):
        raise ValueError("Amount must be a valid number.")

    if amount <= 0:
        raise ValueError("Expense amount must be greater than zero.")

    if not category or not category.strip():
        raise ValueError("Category cannot be empty.")

    if expense_date is None:
        expense_date = dt_date.today().isoformat()
    elif isinstance(expense_date, (datetime, dt_date)):
        expense_date = expense_date.strftime("%Y-%m-%d")

    init_db(db_path)

    with get_connection(db_path) as conn:
        cursor = conn.cursor()
        cursor.execute(
            """
            INSERT INTO expenses (title, amount, category, date)
            VALUES (?, ?, ?, ?)
            """,
            (title.strip(), round(amount, 2), category.strip().capitalize(), str(expense_date)),
        )
        conn.commit()
        return cursor.lastrowid


def get_all_expenses(db_path=DEFAULT_DB):
    """Return a list of all expenses as dictionaries ordered by date and ID."""
    init_db(db_path)
    with get_connection(db_path) as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT id, title, amount, category, date FROM expenses ORDER BY date ASC, id ASC")
        rows = cursor.fetchall()
        return [dict(row) for row in rows]


def get_total_expenses(db_path=DEFAULT_DB):
    """Return the sum of all expenses."""
    init_db(db_path)
    with get_connection(db_path) as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT COALESCE(SUM(amount), 0.0) FROM expenses")
        (total,) = cursor.fetchone()
        return round(float(total), 2)


def get_category_summary(db_path=DEFAULT_DB):
    """Return a dictionary of category -> total amount spent."""
    init_db(db_path)
    with get_connection(db_path) as conn:
        cursor = conn.cursor()
        cursor.execute(
            """
            SELECT category, ROUND(SUM(amount), 2) as total
            FROM expenses
            GROUP BY category
            ORDER BY total DESC
            """
        )
        rows = cursor.fetchall()
        return {row["category"]: row["total"] for row in rows}


def delete_expense(expense_id, db_path=DEFAULT_DB):
    """Delete an expense by ID. Returns True if deleted, False if not found."""
    init_db(db_path)
    with get_connection(db_path) as conn:
        cursor = conn.cursor()
        cursor.execute("DELETE FROM expenses WHERE id = ?", (expense_id,))
        conn.commit()
        return cursor.rowcount > 0


def display_expenses(expenses):
    """Format and print expenses in a clean table format."""
    if not expenses:
        print("\nNo expenses found.")
        return

    print("\n" + "=" * 65)
    print(f"{'ID':<5} {'Date':<12} {'Category':<15} {'Title':<20} {'Amount':>10}")
    print("-" * 65)
    for exp in reversed(expenses):
        print(
            f"{exp['id']:<5} {exp['date']:<12} {exp['category']:<15} "
            f"{exp['title'][:18]:<20} ${exp['amount']:>9.2f}"
        )
    print("=" * 65)


def export_to_excel(db_path=DEFAULT_DB, excel_path=DEFAULT_EXCEL):
    """
    Generate a formatted Excel workbook with:
      - Tab 1 ('Daily Expenses'): For viewing and directly entering daily expenses.
      - Tab 2 ('Summary & Analytics'): Automatically calculates total weekly,
        monthly, and yearly expenses via live Excel formulas.
    """
    if openpyxl is None:
        raise ImportError("openpyxl is required to export to Excel. Run 'pip install openpyxl'.")

    init_db(db_path)
    expenses = get_all_expenses(db_path)

    wb = openpyxl.Workbook()

    # Styling Palette
    NAVY_HEADER = "1B365D"
    WHITE = "FFFFFF"
    BORDER_COLOR = "D9D9D9"
    ACCENT_BLUE = "2A6F97"
    LIGHT_BLUE = "E8F1F5"
    TABLE_HEADER = "335C67"
    TOTAL_ROW_FILL = "E9ECEF"

    thin_border = Border(
        left=Side(style="thin", color=BORDER_COLOR),
        right=Side(style="thin", color=BORDER_COLOR),
        top=Side(style="thin", color=BORDER_COLOR),
        bottom=Side(style="thin", color=BORDER_COLOR),
    )

    # -------------------------------------------------------------
    # 1. TAB 1: Daily Expenses
    # -------------------------------------------------------------
    ws_daily = wb.active
    ws_daily.title = "Daily Expenses"
    ws_daily.views.sheetView[0].showGridLines = True

    headers = [
        "ID",
        "Date (YYYY-MM-DD)",
        "Category",
        "Title / Description",
        "Amount ($)",
        "Week Num",
        "Month",
        "Year",
    ]
    ws_daily.append(headers)

    header_font = Font(name="Calibri", size=11, bold=True, color=WHITE)
    header_fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
    for col_idx in range(1, len(headers) + 1):
        cell = ws_daily.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws_daily.row_dimensions[1].height = 28

    current_row = 2
    for exp in expenses:
        ws_daily.cell(row=current_row, column=1, value=exp["id"]).alignment = Alignment(horizontal="center")

        # Parse date to actual date object if possible
        d_val = exp["date"]
        try:
            parsed_d = datetime.strptime(str(d_val), "%Y-%m-%d").date()
            d_cell = ws_daily.cell(row=current_row, column=2, value=parsed_d)
            d_cell.number_format = "yyyy-mm-dd"
        except (ValueError, TypeError):
            d_cell = ws_daily.cell(row=current_row, column=2, value=str(d_val))
        d_cell.alignment = Alignment(horizontal="center")

        ws_daily.cell(row=current_row, column=3, value=exp["category"]).alignment = Alignment(horizontal="left")
        ws_daily.cell(row=current_row, column=4, value=exp["title"]).alignment = Alignment(horizontal="left")

        amt_cell = ws_daily.cell(row=current_row, column=5, value=float(exp["amount"]))
        amt_cell.number_format = "$#,##0.00"
        amt_cell.alignment = Alignment(horizontal="right")

        # Dynamic cross-compatible formulas for Week, Month, Year
        ws_daily.cell(
            row=current_row, column=6,
            value=f'=IF(B{current_row}="","",ISOWEEKNUM(IF(ISNUMBER(B{current_row}),B{current_row},DATEVALUE(B{current_row}))))'
        ).alignment = Alignment(horizontal="center")

        ws_daily.cell(
            row=current_row, column=7,
            value=f'=IF(B{current_row}="","",TEXT(IF(ISNUMBER(B{current_row}),B{current_row},DATEVALUE(B{current_row})),"yyyy-mm"))'
        ).alignment = Alignment(horizontal="center")

        ws_daily.cell(
            row=current_row, column=8,
            value=f'=IF(B{current_row}="","",YEAR(IF(ISNUMBER(B{current_row}),B{current_row},DATEVALUE(B{current_row}))))'
        ).alignment = Alignment(horizontal="center")

        for c in range(1, 9):
            ws_daily.cell(row=current_row, column=c).border = thin_border
            ws_daily.cell(row=current_row, column=c).font = Font(name="Calibri", size=10)
        current_row += 1

    # Pre-populate empty entry rows with ready formulas so the user can just type in Excel!
    total_setup_rows = max(current_row + 150, 200)
    for r in range(current_row, total_setup_rows + 1):
        d_cell = ws_daily.cell(row=r, column=2)
        d_cell.number_format = "yyyy-mm-dd"

        amt_cell = ws_daily.cell(row=r, column=5)
        amt_cell.number_format = "$#,##0.00"

        ws_daily.cell(
            row=r, column=6,
            value=f'=IF(B{r}="","",ISOWEEKNUM(IF(ISNUMBER(B{r}),B{r},DATEVALUE(B{r}))))'
        ).alignment = Alignment(horizontal="center")

        ws_daily.cell(
            row=r, column=7,
            value=f'=IF(B{r}="","",TEXT(IF(ISNUMBER(B{r}),B{r},DATEVALUE(B{r})),"yyyy-mm"))'
        ).alignment = Alignment(horizontal="center")

        ws_daily.cell(
            row=r, column=8,
            value=f'=IF(B{r}="","",YEAR(IF(ISNUMBER(B{r}),B{r},DATEVALUE(B{r}))))'
        ).alignment = Alignment(horizontal="center")

        for c in range(1, 9):
            ws_daily.cell(row=r, column=c).border = thin_border
            ws_daily.cell(row=r, column=c).font = Font(name="Calibri", size=10)

    # Category Dropdown Validation (Column C)
    dv = DataValidation(
        type="list",
        formula1=f'"{",".join(STANDARD_CATEGORIES)}"',
        allow_blank=True,
        prompt="Select expense category from list",
        promptTitle="Category Selection",
    )
    ws_daily.add_data_validation(dv)
    dv.add(f"C2:C{total_setup_rows}")

    daily_col_widths = {
        "A": 8,   # ID
        "B": 18,  # Date
        "C": 18,  # Category
        "D": 30,  # Description
        "E": 15,  # Amount
        "F": 12,  # Week Num
        "G": 14,  # Month
        "H": 10,  # Year
    }
    for col_letter, width in daily_col_widths.items():
        ws_daily.column_dimensions[col_letter].width = width

    # -------------------------------------------------------------
    # 2. TAB 2: Summary & Analytics (Weekly, Monthly, Yearly)
    # -------------------------------------------------------------
    ws_sum = wb.create_sheet("Summary & Analytics")
    ws_sum.views.sheetView[0].showGridLines = True

    # Title Banner
    ws_sum.merge_cells("A1:G1")
    title_cell = ws_sum["A1"]
    title_cell.value = "📊 PERSONAL EXPENSE TRACKER - SUMMARY & ANALYTICS"
    title_cell.font = Font(name="Calibri", size=15, bold=True, color=WHITE)
    title_cell.fill = header_fill
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_sum.row_dimensions[1].height = 36

    ws_sum.merge_cells("A2:G2")
    sub_cell = ws_sum["A2"]
    sub_cell.value = "Live formulas automatically calculating total weekly, monthly, and yearly expenses as you enter daily records"
    sub_cell.font = Font(name="Calibri", size=10, italic=True, color="555555")
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_sum.row_dimensions[2].height = 20

    # Executive Overview KPI Cards (Rows 4-5)
    kpis = [
        ("B4", "B5", "💰 Total Spending", "=SUM('Daily Expenses'!E2:E1000)", "$#,##0.00"),
        ("C4", "C5", "📝 Total Transactions", "=COUNTA('Daily Expenses'!B2:B1000)", "#,##0"),
        ("D4", "D5", "📅 This Week's Total", "=SUMIFS('Daily Expenses'!E2:E1000, 'Daily Expenses'!F2:F1000, ISOWEEKNUM(TODAY()))", "$#,##0.00"),
        ("E4", "E5", "🗓️ This Month's Total", '=SUMIFS(\'Daily Expenses\'!E2:E1000, \'Daily Expenses\'!G2:G1000, TEXT(TODAY(),"yyyy-mm"))', "$#,##0.00"),
        ("F4", "F5", "📆 This Year's Total", "=SUMIFS('Daily Expenses'!E2:E1000, 'Daily Expenses'!H2:H1000, YEAR(TODAY()))", "$#,##0.00"),
    ]
    kpi_hdr_fill = PatternFill(start_color=ACCENT_BLUE, end_color=ACCENT_BLUE, fill_type="solid")
    kpi_val_fill = PatternFill(start_color=LIGHT_BLUE, end_color=LIGHT_BLUE, fill_type="solid")

    for h_pos, v_pos, label, formula, num_format in kpis:
        c_h = ws_sum[h_pos]
        c_h.value = label
        c_h.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
        c_h.fill = kpi_hdr_fill
        c_h.alignment = Alignment(horizontal="center", vertical="center")
        c_h.border = thin_border

        c_v = ws_sum[v_pos]
        c_v.value = formula
        c_v.font = Font(name="Calibri", size=13, bold=True, color=NAVY_HEADER)
        c_v.fill = kpi_val_fill
        c_v.alignment = Alignment(horizontal="center", vertical="center")
        c_v.number_format = num_format
        c_v.border = thin_border

    ws_sum.row_dimensions[4].height = 24
    ws_sum.row_dimensions[5].height = 28

    tbl_hdr_fill = PatternFill(start_color=TABLE_HEADER, end_color=TABLE_HEADER, fill_type="solid")

    # --- Section: Monthly Expenses (Cols A-C, Row 8+) ---
    ws_sum["A8"] = "🗓️ MONTHLY EXPENSES (2026)"
    ws_sum["A8"].font = Font(name="Calibri", size=12, bold=True, color=NAVY_HEADER)

    month_headers = ["Month Key", "Month Name", "Total Spent ($)"]
    for i, h in enumerate(month_headers, start=1):
        c = ws_sum.cell(row=9, column=i, value=h)
        c.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
        c.fill = tbl_hdr_fill
        c.alignment = Alignment(horizontal="center" if i < 3 else "right")
        c.border = thin_border

    months = [
        ("2026-01", "January 2026"),
        ("2026-02", "February 2026"),
        ("2026-03", "March 2026"),
        ("2026-04", "April 2026"),
        ("2026-05", "May 2026"),
        ("2026-06", "June 2026"),
        ("2026-07", "July 2026"),
        ("2026-08", "August 2026"),
        ("2026-09", "September 2026"),
        ("2026-10", "October 2026"),
        ("2026-11", "November 2026"),
        ("2026-12", "December 2026"),
    ]
    m_row = 10
    for key, name in months:
        ws_sum.cell(row=m_row, column=1, value=key).alignment = Alignment(horizontal="center")
        ws_sum.cell(row=m_row, column=2, value=name).alignment = Alignment(horizontal="left")
        amt_c = ws_sum.cell(
            row=m_row, column=3,
            value=f"=SUMIF('Daily Expenses'!$G$2:$G$1000, A{m_row}, 'Daily Expenses'!$E$2:$E$1000)"
        )
        amt_c.number_format = "$#,##0.00"
        amt_c.alignment = Alignment(horizontal="right")
        for col_i in range(1, 4):
            ws_sum.cell(row=m_row, column=col_i).border = thin_border
            ws_sum.cell(row=m_row, column=col_i).font = Font(name="Calibri", size=10)
        m_row += 1

    # Monthly Total
    ws_sum.cell(row=m_row, column=1, value="Total Year").alignment = Alignment(horizontal="center")
    ws_sum.cell(row=m_row, column=1).font = Font(name="Calibri", size=10, bold=True)
    ws_sum.cell(row=m_row, column=2, value="All Months (2026)").alignment = Alignment(horizontal="left")
    ws_sum.cell(row=m_row, column=2).font = Font(name="Calibri", size=10, bold=True)
    m_tot = ws_sum.cell(row=m_row, column=3, value=f"=SUM(C10:C{m_row-1})")
    m_tot.number_format = "$#,##0.00"
    m_tot.alignment = Alignment(horizontal="right")
    m_tot.font = Font(name="Calibri", size=10, bold=True)
    for col_i in range(1, 4):
        ws_sum.cell(row=m_row, column=col_i).border = thin_border
        ws_sum.cell(row=m_row, column=col_i).fill = PatternFill(start_color=TOTAL_ROW_FILL, end_color=TOTAL_ROW_FILL, fill_type="solid")

    # --- Section: Yearly Expenses (Cols A-C, below Monthly) ---
    y_start = m_row + 2
    ws_sum.cell(row=y_start, column=1, value="📆 YEARLY EXPENSES").font = Font(name="Calibri", size=12, bold=True, color=NAVY_HEADER)

    year_headers = ["Year", "Description", "Total Spent ($)"]
    for i, h in enumerate(year_headers, start=1):
        c = ws_sum.cell(row=y_start + 1, column=i, value=h)
        c.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
        c.fill = tbl_hdr_fill
        c.alignment = Alignment(horizontal="center" if i < 3 else "right")
        c.border = thin_border

    years = [
        (2024, "Past Year (2024)"),
        (2025, "Previous Year (2025)"),
        (2026, "Current Year (2026)"),
        (2027, "Upcoming Year (2027)"),
        (2028, "Future Year (2028)"),
    ]
    y_row = y_start + 2
    for yr, desc in years:
        ws_sum.cell(row=y_row, column=1, value=yr).alignment = Alignment(horizontal="center")
        ws_sum.cell(row=y_row, column=2, value=desc).alignment = Alignment(horizontal="left")
        amt_c = ws_sum.cell(
            row=y_row, column=3,
            value=f"=SUMIF('Daily Expenses'!$H$2:$H$1000, A{y_row}, 'Daily Expenses'!$E$2:$E$1000)"
        )
        amt_c.number_format = "$#,##0.00"
        amt_c.alignment = Alignment(horizontal="right")
        for col_i in range(1, 4):
            ws_sum.cell(row=y_row, column=col_i).border = thin_border
            ws_sum.cell(row=y_row, column=col_i).font = Font(name="Calibri", size=10)
        y_row += 1

    # --- Section: Weekly Expenses (Cols E-G, Row 8+) ---
    ws_sum["E8"] = "📅 WEEKLY EXPENSES (Weeks 1 to 52)"
    ws_sum["E8"].font = Font(name="Calibri", size=12, bold=True, color=NAVY_HEADER)

    weekly_headers = ["Week #", "Week Label", "Total Spent ($)"]
    for i, h in enumerate(weekly_headers, start=5):
        c = ws_sum.cell(row=9, column=i, value=h)
        c.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
        c.fill = tbl_hdr_fill
        c.alignment = Alignment(horizontal="center" if i < 7 else "right")
        c.border = thin_border

    # Build weekly schedule (Weeks 1 to 52)
    w_row = 10
    for wk in range(1, 53):
        ws_sum.cell(row=w_row, column=5, value=wk).alignment = Alignment(horizontal="center")
        ws_sum.cell(row=w_row, column=6, value=f"Week {wk:02d}").alignment = Alignment(horizontal="left")
        amt_c = ws_sum.cell(
            row=w_row, column=7,
            value=f"=SUMIF('Daily Expenses'!$F$2:$F$1000, E{w_row}, 'Daily Expenses'!$E$2:$E$1000)"
        )
        amt_c.number_format = "$#,##0.00"
        amt_c.alignment = Alignment(horizontal="right")
        for col_i in range(5, 8):
            ws_sum.cell(row=w_row, column=col_i).border = thin_border
            ws_sum.cell(row=w_row, column=col_i).font = Font(name="Calibri", size=10)
        w_row += 1

    # --- Section: Category Summary (Cols I-K or below Yearly) ---
    cat_start = y_row + 2
    ws_sum.cell(row=cat_start, column=1, value="🏷️ SPENDING BY CATEGORY").font = Font(name="Calibri", size=12, bold=True, color=NAVY_HEADER)

    cat_headers = ["Category", "% Share", "Total Spent ($)"]
    for i, h in enumerate(cat_headers, start=1):
        c = ws_sum.cell(row=cat_start + 1, column=i, value=h)
        c.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
        c.fill = tbl_hdr_fill
        c.alignment = Alignment(horizontal="center" if i == 2 else ("right" if i == 3 else "left"))
        c.border = thin_border

    cat_row = cat_start + 2
    for cat in STANDARD_CATEGORIES:
        ws_sum.cell(row=cat_row, column=1, value=cat).alignment = Alignment(horizontal="left")

        share_c = ws_sum.cell(row=cat_row, column=2, value=f'=IF($B$5>0, C{cat_row}/$B$5, 0)')
        share_c.number_format = "0.0%"
        share_c.alignment = Alignment(horizontal="center")

        amt_c = ws_sum.cell(
            row=cat_row, column=3,
            value=f"=SUMIF('Daily Expenses'!$C$2:$C$1000, A{cat_row}, 'Daily Expenses'!$E$2:$E$1000)"
        )
        amt_c.number_format = "$#,##0.00"
        amt_c.alignment = Alignment(horizontal="right")

        for col_i in range(1, 4):
            ws_sum.cell(row=cat_row, column=col_i).border = thin_border
            ws_sum.cell(row=cat_row, column=col_i).font = Font(name="Calibri", size=10)
        cat_row += 1

    sum_col_widths = {
        "A": 16,
        "B": 22,
        "C": 20,
        "D": 8,
        "E": 12,
        "F": 18,
        "G": 18,
    }
    for col_letter, width in sum_col_widths.items():
        ws_sum.column_dimensions[col_letter].width = width

    wb.save(excel_path)
    return excel_path


def import_from_excel(excel_path=DEFAULT_EXCEL, db_path=DEFAULT_DB):
    """
    Read newly added rows from the 'Daily Expenses' sheet of the Excel file
    and insert them into SQLite database if not already present.
    Returns the number of new records imported.
    """
    if openpyxl is None:
        raise ImportError("openpyxl is required to import from Excel. Run 'pip install openpyxl'.")

    if not os.path.exists(excel_path):
        raise FileNotFoundError(f"Excel file not found at: {excel_path}")

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    if "Daily Expenses" not in wb.sheetnames:
        raise ValueError("Workbook must contain a 'Daily Expenses' sheet.")

    ws = wb["Daily Expenses"]
    existing = get_all_expenses(db_path)

    # Set of existing expense signatures: (date, title_lower, amount, category_lower)
    existing_sigs = {
        (str(e["date"]), e["title"].strip().lower(), round(float(e["amount"]), 2), e["category"].strip().lower())
        for e in existing
    }

    imported_count = 0
    # Data starts from row 2
    for r in range(2, ws.max_row + 1):
        date_val = ws.cell(row=r, column=2).value
        cat_val = ws.cell(row=r, column=3).value
        title_val = ws.cell(row=r, column=4).value
        amt_val = ws.cell(row=r, column=5).value

        # Skip rows with no title or amount or date
        if not title_val or amt_val is None or not date_val:
            continue

        # Format date value
        if isinstance(date_val, (datetime, dt_date)):
            date_str = date_val.strftime("%Y-%m-%d")
        else:
            date_str = str(date_val).strip()

        try:
            amount = float(amt_val)
        except (ValueError, TypeError):
            continue

        if amount <= 0:
            continue

        cat_str = str(cat_val).strip() if cat_val else "Other"
        title_str = str(title_val).strip()

        sig = (date_str, title_str.lower(), round(amount, 2), cat_str.lower())
        if sig not in existing_sigs:
            add_expense(title_str, amount, cat_str, date_str, db_path)
            existing_sigs.add(sig)
            imported_count += 1

    return imported_count


def inspect_excel(excel_path=DEFAULT_EXCEL, db_path=DEFAULT_DB):
    """
    Print an inspection of the Excel workbook in terminal,
    showing both worksheets, data rows, formulas, and calculated totals.
    """
    if openpyxl is None:
        raise ImportError("openpyxl is required to view Excel. Run 'pip install openpyxl'.")

    if not os.path.exists(excel_path):
        print(f"Excel file '{excel_path}' not found. Generating now...")
        export_to_excel(db_path, excel_path)

    wb = openpyxl.load_workbook(excel_path, data_only=False)

    print("\n" + "=" * 75)
    print(f" 📑 EXCEL WORKBOOK INSPECTION: {excel_path}")
    print("=" * 75)
    print(f"Sheets Found in Workbook: {wb.sheetnames}\n")

    # 1. Daily Expenses
    ws1 = wb["Daily Expenses"]
    print("━" * 75)
    print(f" ▶ TAB 1: [{ws1.title}] - Daily Expense Records")
    print("━" * 75)
    print(f"{'ID':<4} {'Date':<12} {'Category':<14} {'Title':<20} {'Amount':>10}")
    print("-" * 75)

    count = 0
    for r in range(2, ws1.max_row + 1):
        id_val = ws1.cell(row=r, column=1).value
        date_val = ws1.cell(row=r, column=2).value
        cat_val = ws1.cell(row=r, column=3).value
        title_val = ws1.cell(row=r, column=4).value
        amt_val = ws1.cell(row=r, column=5).value

        if not title_val and amt_val is None and not date_val:
            continue

        count += 1
        d_str = date_val.strftime("%Y-%m-%d") if isinstance(date_val, (datetime, dt_date)) else str(date_val)
        amt_str = f"${float(amt_val):>9.2f}" if amt_val is not None else "$    0.00"
        print(f"{str(id_val or ''):<4} {d_str:<12} {str(cat_val or ''):<14} {str(title_val or '')[:18]:<20} {amt_str}")

    print("-" * 75)
    print(f"Active Records: {count} | Pre-configured entry rows with formulas ready: ~150")

    # 2. Summary & Analytics
    ws2 = wb["Summary & Analytics"]
    print("\n" + "━" * 75)
    print(f" ▶ TAB 2: [{ws2.title}] - Calculated Summaries & Formulas")
    print("━" * 75)

    tot = get_total_expenses(db_path)
    cat_summary = get_category_summary(db_path)

    print("📊 Real-Time KPI Cards (Live Excel Formulas):")
    print(f"  • Total Spending       : ${tot:.2f}  [Formula: {ws2['B5'].value}]")
    print(f"  • Total Transactions   : {count} transactions  [Formula: {ws2['C5'].value}]")
    print(f"  • This Week's Spending : Formula: {ws2['D5'].value}")
    print(f"  • This Month's Spending: Formula: {ws2['E5'].value}")
    print(f"  • This Year's Spending : Formula: {ws2['F5'].value}")

    print("\n🗓️ Monthly Summary Table (Formulas):")
    print(f"  • Month: {ws2['A10'].value} ({ws2['B10'].value}) -> Formula: {ws2['C10'].value}")
    print(f"  • Month: {ws2['A19'].value} ({ws2['B19'].value}) -> Formula: {ws2['C19'].value}")

    print("\n📅 Weekly Summary Schedule (Weeks 1 to 52):")
    print(f"  • Sample Week 01 -> Formula: {ws2['G10'].value}")
    print(f"  • Sample Week 41 -> Formula: {ws2['G50'].value}")

    print("\n📆 Yearly Summary Table (Formulas):")
    print(f"  • Year 2026 -> Formula: {ws2['C28'].value}")

    print("\n🏷️ Category Spending Breakdown:")
    for cat, amt in cat_summary.items():
        pct = (amt / tot * 100) if tot > 0 else 0
        bar = "█" * int(pct // 10)
        print(f"  • {cat:<14}: ${amt:>7.2f} ({pct:>5.1f}%) {bar}")
    print("=" * 75 + "\n")


def run_demo(db_path=DEFAULT_DB, excel_path=DEFAULT_EXCEL):
    """Run an automated demonstration of the expense tracker and generate Excel file."""
    print("=" * 60)
    print("   PERSONAL EXPENSE TRACKER - DEMO")
    print("=" * 60)

    init_db(db_path)

    # Add sample expenses if empty
    if not get_all_expenses(db_path):
        add_expense("Groceries", 45.50, "Food", "2026-10-10", db_path)
        add_expense("Metro Pass", 20.00, "Transport", "2026-10-11", db_path)
        add_expense("Coffee with Friends", 12.75, "Food", "2026-10-12", db_path)
        add_expense("Internet Bill", 55.00, "Bills", "2026-10-13", db_path)
        print(">> Added initial demo expenses.")

    expenses = get_all_expenses(db_path)
    display_expenses(expenses)

    total = get_total_expenses(db_path)
    print(f"\n💰 Total Spending: ${total:.2f}")

    summary = get_category_summary(db_path)
    print("\n📊 Spending by Category:")
    for cat, amt in summary.items():
        percentage = (amt / total * 100) if total > 0 else 0
        bar = "█" * int(percentage // 10)
        print(f"  {cat:<12}: ${amt:>7.2f} ({percentage:>5.1f}%) {bar}")
    print("=" * 60)

    # Export to Excel
    export_to_excel(db_path, excel_path)
    print(f"📁 Excel Document Created/Updated: {excel_path}")
    print("   ├── Tab 1: 'Daily Expenses'       -> Enter daily expenses with dropdowns")
    print("   └── Tab 2: 'Summary & Analytics'  -> Live Weekly, Monthly & Yearly totals")
    print("=" * 60)


def main():
    """Command-line interface entry point."""
    args = sys.argv[1:]

    if "--view-excel" in args:
        idx = args.index("--view-excel")
        file_path = args[idx + 1] if idx + 1 < len(args) else DEFAULT_EXCEL
        inspect_excel(file_path, DEFAULT_DB)
    elif "--export-excel" in args:
        idx = args.index("--export-excel")
        out_path = args[idx + 1] if idx + 1 < len(args) else DEFAULT_EXCEL
        export_to_excel(DEFAULT_DB, out_path)
        print(f"✅ Successfully exported expenses to: {out_path}")
    elif "--import-excel" in args:
        idx = args.index("--import-excel")
        in_path = args[idx + 1] if idx + 1 < len(args) else DEFAULT_EXCEL
        count = import_from_excel(in_path, DEFAULT_DB)
        print(f"✅ Imported {count} new expense(s) from: {in_path}")
    else:
        run_demo()


if __name__ == "__main__":
    main()

